"""FDD workbook extraction, validation, and internal reconciliation.

The workbook is the source of truth.  This module deliberately finds values
from section titles, row labels, and year/header labels instead of treating
cell coordinates as a public data contract.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
from math import isfinite
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook


FDD_SHEET_NAME = "FDD 조정"
FORECAST_YEARS = (2026, 2027, 2028, 2029, 2030)
RECONCILIATION_TOLERANCE = 1e-6
EXCEL_ERROR_PREFIX = "#"


@dataclass(frozen=True)
class QoEReview:
    reported_ebitda: float
    fdd_ebitda: float
    adjustments_2025: Mapping[str, float]
    forecast_ebit_adjustment: Mapping[int, float]
    forecast_da_adjustment: Mapping[int, float]
    lease_capex: Mapping[int, float]
    lease_capex_ratio: Mapping[int, float]


@dataclass(frozen=True)
class NWCReview:
    normalized_peg: float
    closing_nwc: float
    nwc_gap: float
    price_adjustment_recognition_rate: float
    applied_nwc_price_adjustment: float
    capex_payable_reclassification: float
    other_operating_current_liability_ratio: Mapping[int, float]


@dataclass(frozen=True)
class TransactionBridgeReview:
    cash_like: float
    debt_like: float
    net_debt: float
    net_cash: float
    non_operating_assets: float
    non_controlling_interests: float
    applied_nwc_price_adjustment: float
    fdd_equity_adjustment: float
    cash_like_components: Mapping[str, float]
    debt_like_components: Mapping[str, float]
    non_operating_asset_components: Mapping[str, float]
    recognition_rates: Mapping[str, float]


@dataclass(frozen=True)
class FDDInputs:
    qoe: QoEReview
    nwc: NWCReview
    transaction_bridge: TransactionBridgeReview
    reconciliation_differences: Mapping[str, float]

    def to_model_dict(self) -> dict[str, object]:
        """Return a dashboard-safe dictionary without exposing dataclasses."""

        return asdict(self)


def _normalized(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.split())


def _find_unique_row(
    worksheet,
    label: str,
    *,
    start_row: int = 1,
    end_row: int | None = None,
    label_column: int = 2,
) -> int:
    normalized_label = _normalized(label)
    last_row = end_row or worksheet.max_row
    matches = [
        row
        for row in range(start_row, last_row + 1)
        if _normalized(worksheet.cell(row=row, column=label_column).value)
        == normalized_label
    ]
    if not matches:
        raise KeyError(
            f"{worksheet.title!r} 시트에서 필수 라벨 {label!r}을 찾지 못했습니다."
        )
    if len(matches) != 1:
        raise ValueError(
            f"{worksheet.title!r} 시트의 라벨 {label!r}은(는) "
            f"정확히 1개여야 합니다. 현재 {len(matches)}개입니다."
        )
    return matches[0]


def _section_bounds(worksheet, title: str, next_title: str | None) -> tuple[int, int]:
    start = _find_unique_row(worksheet, title)
    end = worksheet.max_row
    if next_title is not None:
        next_start = _find_unique_row(worksheet, next_title)
        if next_start <= start:
            raise ValueError(f"FDD 섹션 순서가 올바르지 않습니다: {title!r}")
        end = next_start - 1
    return start, end


def _header_columns(
    worksheet,
    *,
    section: tuple[int, int],
    first_header: str = "항목",
) -> tuple[int, dict[str, int]]:
    header_row = _find_unique_row(
        worksheet,
        first_header,
        start_row=section[0],
        end_row=section[1],
    )
    headers: dict[str, int] = {}
    for cell in worksheet[header_row]:
        key = _normalized(cell.value)
        if not key:
            continue
        if key in headers:
            raise ValueError(
                f"{worksheet.title!r} 시트 {header_row}행의 header {key!r}가 중복됩니다."
            )
        headers[key] = cell.column
    return header_row, headers


def _year_from_header(value: object) -> int | None:
    if isinstance(value, (datetime, date)):
        return value.year
    if isinstance(value, int) and 1900 <= value <= 2200:
        return value
    normalized = _normalized(value)
    if len(normalized) >= 4 and normalized[:4].isdigit():
        return int(normalized[:4])
    return None


def _labeled_year_series(
    formula_sheet,
    value_sheet,
    label: str,
    years: tuple[int, ...],
    *,
    header_label: str,
) -> dict[int, float]:
    header_row = _find_unique_row(formula_sheet, header_label)
    year_columns: dict[int, int] = {}
    for cell in formula_sheet[header_row]:
        year = _year_from_header(cell.value)
        if year is None:
            continue
        if year in year_columns:
            raise ValueError(
                f"{formula_sheet.title!r} 시트에서 {year}년 header가 중복됩니다."
            )
        year_columns[year] = cell.column
    missing_years = [year for year in years if year not in year_columns]
    if missing_years:
        raise KeyError(
            f"{formula_sheet.title!r} 시트에 연도 header가 없습니다: {missing_years}"
        )
    row = _find_unique_row(formula_sheet, label)
    return {
        year: _cached_value(
            formula_sheet,
            value_sheet,
            row,
            year_columns[year],
            f"{label} / {year}",
        )
        for year in years
    }


def _finite_number(value: object, context: str) -> float:
    if value is None:
        raise ValueError(f"{context}의 cached value가 없습니다.")
    if isinstance(value, str) and value.startswith(EXCEL_ERROR_PREFIX):
        raise ValueError(f"{context}에 Excel 오류 {value!r}가 있습니다.")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TypeError(f"{context}은(는) 숫자여야 합니다. 현재 값: {value!r}")
    result = float(value)
    if not isfinite(result):
        raise ValueError(f"{context}은(는) 유한한 숫자여야 합니다.")
    return result


def validate_recognition_rate(value: object, context: str) -> float:
    rate = _finite_number(value, context)
    if not 0 <= rate <= 1:
        raise ValueError(f"{context}은(는) 0~100% 범위여야 합니다. 현재 값: {rate}")
    return rate


def _cached_value(formula_sheet, value_sheet, row: int, column: int, context: str) -> float:
    formula_cell = formula_sheet.cell(row=row, column=column)
    value_cell = value_sheet.cell(row=row, column=column)
    value = value_cell.value
    if formula_cell.data_type == "f" and value is None:
        raise ValueError(
            f"{context} ({formula_sheet.title}!{formula_cell.coordinate})의 "
            "formula cached value가 없습니다."
        )
    return _finite_number(value, f"{context} ({formula_sheet.title}!{formula_cell.coordinate})")


def _labeled_value(
    formula_sheet,
    value_sheet,
    label: str,
    header: str,
    *,
    section: tuple[int, int],
) -> float:
    _, headers = _header_columns(formula_sheet, section=section)
    normalized_header = _normalized(header)
    if normalized_header not in headers:
        raise KeyError(
            f"{formula_sheet.title!r} 시트에서 header {header!r}을 찾지 못했습니다."
        )
    row = _find_unique_row(
        formula_sheet,
        label,
        start_row=section[0],
        end_row=section[1],
    )
    return _cached_value(
        formula_sheet,
        value_sheet,
        row,
        headers[normalized_header],
        f"{label} / {header}",
    )


def _component_block(
    formula_sheet,
    value_sheet,
    labels: tuple[str, ...],
    *,
    section: tuple[int, int],
) -> tuple[dict[str, float], dict[str, float]]:
    values: dict[str, float] = {}
    rates: dict[str, float] = {}
    for label in labels:
        values[label] = _labeled_value(
            formula_sheet,
            value_sheet,
            label,
            "FDD 반영액",
            section=section,
        )
        rate = _labeled_value(
            formula_sheet,
            value_sheet,
            label,
            "인정률",
            section=section,
        )
        rates[label] = validate_recognition_rate(rate, f"{label} 인정률")
    return values, rates


def _assert_reconciled(name: str, calculated: float, workbook_value: float) -> float:
    difference = calculated - workbook_value
    if abs(difference) > RECONCILIATION_TOLERANCE:
        raise ValueError(
            f"FDD 대사 실패 - {name}: 재계산={calculated:,.6f}, "
            f"workbook={workbook_value:,.6f}, 차이={difference:,.6f}"
        )
    return difference


def _scan_fdd_formula_cache(formula_sheet, value_sheet) -> None:
    issues = []
    for row in formula_sheet.iter_rows():
        for cell in row:
            if cell.data_type != "f":
                continue
            cached = value_sheet[cell.coordinate].value
            if cached is None or (
                isinstance(cached, str) and cached.startswith(EXCEL_ERROR_PREFIX)
            ):
                issues.append(f"{formula_sheet.title}!{cell.coordinate}={cached!r}")
    if issues:
        raise ValueError("FDD formula cache 오류: " + ", ".join(issues[:10]))


def extract_fdd_inputs(excel_path: str | Path) -> FDDInputs:
    """Extract, validate, and reconcile the FDD source-of-truth workbook."""

    path = Path(excel_path)
    formula_workbook = load_workbook(path, data_only=False, read_only=True)
    value_workbook = load_workbook(path, data_only=True, read_only=True)
    if FDD_SHEET_NAME not in formula_workbook.sheetnames:
        raise KeyError(f"필수 시트 {FDD_SHEET_NAME!r}가 없습니다.")

    formula_sheet = formula_workbook[FDD_SHEET_NAME]
    value_sheet = value_workbook[FDD_SHEET_NAME]
    _scan_fdd_formula_cache(formula_sheet, value_sheet)

    qoe_section = _section_bounds(
        formula_sheet,
        "Quality of Earnings 및 지속가능 EBITDA",
        "정상 운전자본 및 Purchase Price Adjustment",
    )
    nwc_section = _section_bounds(
        formula_sheet,
        "정상 운전자본 및 Purchase Price Adjustment",
        "순차입금 및 Debt-like 검토",
    )
    debt_section = _section_bounds(
        formula_sheet,
        "순차입금 및 Debt-like 검토",
        "비영업자산 및 지분가치 조정",
    )
    non_operating_section = _section_bounds(
        formula_sheet,
        "비영업자산 및 지분가치 조정",
        None,
    )

    qoe_adjustment_labels = (
        "비영업 임대수익 재분류",
        "회계·분류 조정",
        "비경상 수익 차감",
        "비경상 비용 가산",
        "특수관계자·보수 정상화",
        "Run-rate / Pro forma 조정",
    )
    reported_ebitda = _labeled_value(
        formula_sheet, value_sheet, "공시 EBITDA", "2025A", section=qoe_section
    )
    qoe_adjustments = {
        label: _labeled_value(
            formula_sheet, value_sheet, label, "2025A", section=qoe_section
        )
        for label in qoe_adjustment_labels
    }
    fdd_ebitda = _labeled_value(
        formula_sheet, value_sheet, "FDD 조정 EBITDA", "2025A", section=qoe_section
    )

    forecast_ebit_adjustment = {}
    forecast_da_adjustment = {}
    lease_capex = {}
    lease_capex_ratio = {}
    for year in FORECAST_YEARS:
        year_header = f"{year}E"
        forecast_ebit_adjustment[year] = _labeled_value(
            formula_sheet,
            value_sheet,
            "투자부동산 EBIT 재분류",
            year_header,
            section=qoe_section,
        )
        forecast_da_adjustment[year] = _labeled_value(
            formula_sheet,
            value_sheet,
            "투자부동산 D&A 조정",
            year_header,
            section=qoe_section,
        )
        lease_capex[year] = _labeled_value(
            formula_sheet,
            value_sheet,
            "사용권자산 추가(Lease capex)",
            year_header,
            section=qoe_section,
        )
        lease_capex_ratio[year] = validate_recognition_rate(
            _labeled_value(
                formula_sheet,
                value_sheet,
                "Lease capex / 매출액",
                year_header,
                section=qoe_section,
            ),
            f"{year}E Lease capex / 매출액",
        )

    normalized_peg = _labeled_value(
        formula_sheet, value_sheet, "2025년 정상 NWC Peg", "2025A", section=nwc_section
    )
    closing_nwc = _labeled_value(
        formula_sheet, value_sheet, "2025년 Closing NWC", "2025A", section=nwc_section
    )
    nwc_gap = _labeled_value(
        formula_sheet, value_sheet, "NWC 초과 / (부족)", "2025A", section=nwc_section
    )
    nwc_rate = validate_recognition_rate(
        _labeled_value(
            formula_sheet,
            value_sheet,
            "가격조정 적용률(0~100%)",
            "2025A",
            section=nwc_section,
        ),
        "NWC 가격조정 적용률",
    )
    applied_nwc = _labeled_value(
        formula_sheet, value_sheet, "적용 NWC 가격조정", "2025A", section=nwc_section
    )
    capex_payable = _labeled_value(
        formula_sheet,
        value_sheet,
        "NWC 내 debt-like 재분류",
        "2025A",
        section=nwc_section,
    )
    revenue_2025 = _labeled_value(
        formula_sheet, value_sheet, "매출액", "2025A", section=nwc_section
    )
    if revenue_2025 <= 0:
        raise ValueError("2025A 매출액은 0보다 커야 합니다.")
    assumptions_formula = formula_workbook["가정"]
    assumptions_value = value_workbook["가정"]
    other_liability_ratio = {
        year: validate_recognition_rate(value, f"{year}E 기타 영업유동부채 비율")
        for year, value in _labeled_year_series(
            assumptions_formula,
            assumptions_value,
            "기타 영업유동부채 / 매출액",
            FORECAST_YEARS,
            header_label="동인",
        ).items()
    }

    cash_labels = (
        "현금및현금성자산",
        "필요 영업현금",
        "단기금융상품",
        "유동 FVTPL 금융자산",
        "기타유동금융자산",
    )
    debt_labels = (
        "금융기관차입금",
        "리스부채",
        "공급자금융약정 대상 채무",
        "Capex 관련 미지급금",
        "분할 관련 연대채무",
        "소송 총 청구액",
        "기타 debt-like (당기법인세부채 등)",
    )
    non_operating_labels = (
        "리가켐바이오 시장가치",
        "기타 관계·공동기업",
        "투자부동산 공정가치",
        "비유동 OCI 금융자산",
        "순확정급여자산",
        "기타 비영업자산",
    )
    cash_components, cash_rates = _component_block(
        formula_sheet, value_sheet, cash_labels, section=debt_section
    )
    debt_components, debt_rates = _component_block(
        formula_sheet, value_sheet, debt_labels, section=debt_section
    )
    non_operating_components, non_operating_rates = _component_block(
        formula_sheet,
        value_sheet,
        non_operating_labels,
        section=non_operating_section,
    )
    nci = _labeled_value(
        formula_sheet,
        value_sheet,
        "비지배지분",
        "FDD 반영액",
        section=non_operating_section,
    )
    nci_rate = validate_recognition_rate(
        _labeled_value(
            formula_sheet,
            value_sheet,
            "비지배지분",
            "인정률",
            section=non_operating_section,
        ),
        "비지배지분 인정률",
    )

    cash_like = _labeled_value(
        formula_sheet, value_sheet, "Cash-like 자산", "FDD 반영액", section=debt_section
    )
    debt_like = _labeled_value(
        formula_sheet, value_sheet, "Debt-like 항목", "FDD 반영액", section=debt_section
    )
    net_debt = _labeled_value(
        formula_sheet,
        value_sheet,
        "순차입금 / (순현금)",
        "FDD 반영액",
        section=debt_section,
    )
    non_operating_assets = _labeled_value(
        formula_sheet,
        value_sheet,
        "비영업자산 합계",
        "FDD 반영액",
        section=non_operating_section,
    )
    applied_nwc_bridge = _labeled_value(
        formula_sheet,
        value_sheet,
        "적용 NWC 가격조정",
        "FDD 반영액",
        section=non_operating_section,
    )
    fdd_equity_adjustment = _labeled_value(
        formula_sheet,
        value_sheet,
        "FDD 지분가치 조정액",
        "FDD 반영액",
        section=non_operating_section,
    )

    calculated_fdd_ebitda = reported_ebitda + sum(qoe_adjustments.values())
    calculated_nwc_gap = closing_nwc - normalized_peg
    calculated_applied_nwc = calculated_nwc_gap * nwc_rate
    calculated_cash_like = sum(cash_components.values())
    calculated_debt_like = sum(debt_components.values())
    calculated_net_debt = calculated_debt_like - calculated_cash_like
    calculated_non_operating_assets = sum(non_operating_components.values())
    calculated_equity_adjustment = (
        -calculated_net_debt
        + calculated_non_operating_assets
        - nci
        + applied_nwc_bridge
    )

    differences = {
        "FDD EBITDA": _assert_reconciled(
            "FDD EBITDA", calculated_fdd_ebitda, fdd_ebitda
        ),
        "NWC gap": _assert_reconciled("NWC gap", calculated_nwc_gap, nwc_gap),
        "Applied NWC PPA": _assert_reconciled(
            "Applied NWC PPA", calculated_applied_nwc, applied_nwc
        ),
        "Applied NWC bridge": _assert_reconciled(
            "Applied NWC bridge", applied_nwc, applied_nwc_bridge
        ),
        "Cash-like": _assert_reconciled(
            "Cash-like", calculated_cash_like, cash_like
        ),
        "Debt-like": _assert_reconciled(
            "Debt-like", calculated_debt_like, debt_like
        ),
        "Net debt": _assert_reconciled(
            "Net debt", calculated_net_debt, net_debt
        ),
        "Non-operating assets": _assert_reconciled(
            "Non-operating assets",
            calculated_non_operating_assets,
            non_operating_assets,
        ),
        "FDD equity adjustment": _assert_reconciled(
            "FDD equity adjustment",
            calculated_equity_adjustment,
            fdd_equity_adjustment,
        ),
    }

    recognition_rates = {
        **{f"Cash-like / {key}": value for key, value in cash_rates.items()},
        **{f"Debt-like / {key}": value for key, value in debt_rates.items()},
        **{
            f"Non-operating assets / {key}": value
            for key, value in non_operating_rates.items()
        },
        "NCI": nci_rate,
        "Applied NWC PPA": nwc_rate,
    }

    return FDDInputs(
        qoe=QoEReview(
            reported_ebitda=reported_ebitda,
            fdd_ebitda=fdd_ebitda,
            adjustments_2025=qoe_adjustments,
            forecast_ebit_adjustment=forecast_ebit_adjustment,
            forecast_da_adjustment=forecast_da_adjustment,
            lease_capex=lease_capex,
            lease_capex_ratio=lease_capex_ratio,
        ),
        nwc=NWCReview(
            normalized_peg=normalized_peg,
            closing_nwc=closing_nwc,
            nwc_gap=nwc_gap,
            price_adjustment_recognition_rate=nwc_rate,
            applied_nwc_price_adjustment=applied_nwc,
            capex_payable_reclassification=capex_payable,
            other_operating_current_liability_ratio=other_liability_ratio,
        ),
        transaction_bridge=TransactionBridgeReview(
            cash_like=cash_like,
            debt_like=debt_like,
            net_debt=net_debt,
            net_cash=-net_debt,
            non_operating_assets=non_operating_assets,
            non_controlling_interests=nci,
            applied_nwc_price_adjustment=applied_nwc_bridge,
            fdd_equity_adjustment=fdd_equity_adjustment,
            cash_like_components=cash_components,
            debt_like_components=debt_components,
            non_operating_asset_components=non_operating_components,
            recognition_rates=recognition_rates,
        ),
        reconciliation_differences=differences,
    )
