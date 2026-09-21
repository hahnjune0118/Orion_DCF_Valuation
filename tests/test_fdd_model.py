from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from fdd_model import (
    _find_unique_row,
    _finite_number,
    extract_fdd_inputs,
    validate_recognition_rate,
)
from orion_dcf import run_orion_dcf


ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "raw" / "orion_dcf.xlsx"


@pytest.fixture(scope="module")
def fdd():
    return extract_fdd_inputs(EXCEL_PATH)


@pytest.fixture(scope="module")
def model():
    return run_orion_dcf(EXCEL_PATH)


def test_fdd_headline_inputs_reconcile(fdd):
    assert fdd.qoe.reported_ebitda == pytest.approx(722_873.922, abs=0.001)
    assert fdd.qoe.fdd_ebitda == pytest.approx(718_602.217, abs=0.001)
    assert fdd.nwc.normalized_peg == pytest.approx(117_489.415, abs=0.001)
    assert fdd.nwc.closing_nwc == pytest.approx(116_509.367, abs=0.001)
    assert fdd.nwc.nwc_gap == pytest.approx(-980.048, abs=0.001)
    assert fdd.nwc.applied_nwc_price_adjustment == pytest.approx(0.0)
    assert fdd.nwc.capex_payable_reclassification == pytest.approx(13_687.0)
    assert all(
        value == pytest.approx(0.0809222964080706)
        for value in fdd.nwc.other_operating_current_liability_ratio.values()
    )
    assert all(abs(value) <= 1e-6 for value in fdd.reconciliation_differences.values())


def test_fdd_dashboard_register_and_forecast_inputs_are_exposed(fdd):
    assert fdd.qoe.adjustments_2025["비영업 임대수익 재분류"] == pytest.approx(
        -4_271.705, abs=0.001
    )
    assert fdd.qoe.forecast_ebit_adjustment[2026] == pytest.approx(
        -4_055.074, abs=0.001
    )
    assert fdd.qoe.forecast_da_adjustment[2026] == pytest.approx(
        -216.631, abs=0.001
    )
    assert fdd.qoe.lease_capex_ratio[2030] == pytest.approx(0.0038)

    bridge = fdd.transaction_bridge
    assert bridge.cash_like_components["단기금융상품"] == pytest.approx(
        986_075.623, abs=0.001
    )
    assert bridge.debt_like_components["리스부채"] == pytest.approx(
        35_516.035, abs=0.001
    )
    assert bridge.debt_like_components["Capex 관련 미지급금"] == pytest.approx(
        13_687.0, abs=0.001
    )
    assert bridge.non_operating_asset_components[
        "리가켐바이오 시장가치"
    ] == pytest.approx(1_626_402.257, abs=0.001)
    assert bridge.recognition_rates

def test_fdd_transaction_bridge_reconciles(fdd):
    bridge = fdd.transaction_bridge
    assert bridge.cash_like == pytest.approx(1_235_817.717, abs=0.001)
    assert bridge.debt_like == pytest.approx(49_203.035, abs=0.001)
    assert bridge.net_cash == pytest.approx(1_186_614.682, abs=0.001)
    assert bridge.non_operating_assets == pytest.approx(1_801_499.020, abs=0.001)
    assert bridge.non_controlling_interests == pytest.approx(103_543.703, abs=0.001)
    assert bridge.fdd_equity_adjustment == pytest.approx(2_884_569.999, abs=0.001)


def test_fdd_base_case_reconciles_to_excel(model):
    expected_fcff = [
        269_051.075,
        318_693.548,
        532_375.101,
        535_014.336,
        562_920.607,
    ]
    assert [row["FCFF"] for row in model["전망"]] == pytest.approx(
        expected_fcff, abs=0.001
    )
    assert model["DCF"]["기업가치"] == pytest.approx(6_530_454.168, abs=0.001)
    assert model["지분가치"]["지분가치"] == pytest.approx(
        9_415_024.166, abs=0.001
    )
    assert model["지분가치"]["주당 내재가치"] == pytest.approx(
        238_181.453, abs=0.001
    )


def test_forecast_rows_expose_fdd_adjusted_values(model):
    for row in model["전망"]:
        assert row["EBIT"] == pytest.approx(row["FDD 조정 EBIT"])
        assert row["D&A"] == pytest.approx(row["FDD 조정 D&A"])
        assert row["Capex"] == pytest.approx(row["FDD 조정 Capex"])
        assert row["NWC"] == pytest.approx(row["FDD 조정 NWC"])
        assert row["NWC 증감"] == pytest.approx(row["FDD 조정 NWC 증감"])
        assert row["FCFF"] == pytest.approx(row["FDD 조정 FCFF"])
        assert row["FDD EBIT 조정액"] == pytest.approx(-4_055.074)
        assert row["FDD D&A 조정액"] == pytest.approx(-216.631)
        assert row["Lease capex"] == pytest.approx(row["매출액"] * 0.0038)


def test_missing_fdd_sheet_fails_fast(tmp_path):
    path = tmp_path / "missing_fdd.xlsx"
    Workbook().save(path)
    with pytest.raises(KeyError, match="FDD 조정"):
        extract_fdd_inputs(path)


def test_missing_formula_cache_fails_fast(tmp_path):
    path = tmp_path / "missing_cache.xlsx"
    workbook = load_workbook(EXCEL_PATH, data_only=False)
    workbook.save(path)
    with pytest.raises(ValueError, match="formula cache"):
        extract_fdd_inputs(path)


def test_duplicate_and_missing_labels_fail_fast():
    worksheet = Workbook().active
    worksheet["B1"] = "중복"
    worksheet["B2"] = "중복"
    with pytest.raises(ValueError, match="정확히 1개"):
        _find_unique_row(worksheet, "중복")
    with pytest.raises(KeyError, match="찾지 못했습니다"):
        _find_unique_row(worksheet, "누락")


@pytest.mark.parametrize("value", [-0.01, 1.01, float("inf"), float("nan")])
def test_invalid_recognition_rates_fail_fast(value):
    with pytest.raises(ValueError):
        validate_recognition_rate(value, "테스트 인정률")


@pytest.mark.parametrize("value", [None, "#NAME?", "#VALUE!", "#REF!"])
def test_missing_or_error_values_fail_fast(value):
    with pytest.raises((TypeError, ValueError)):
        _finite_number(value, "필수 FDD 값")
